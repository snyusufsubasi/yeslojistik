#!/usr/bin/env python3
"""pratikortam dışa aktarımlarını panelin Excel aktarım şablonlarına çevirir.

Kullanım:  python3 tools/legacy/transform.py <extract çıktı klasörü>   (gerekli: pip install openpyxl)
Girdi:     <klasör>/exports/*.html|xlsx  (extract.mjs)
Çıktı:     <klasör>/aktar/1-tedarikciler.xlsx … 10-personeller.xlsx  +  rapor.txt   (repoya girmez)

Kurallar (kullanıcı kararı, 2 Ekim):
  * Yalnız kayıtlar taşınır, borç/alacak çıkarılmaz: müşteri ve tedarikçi devri 0, bütün seferler ESKİ KAYIT
    (taşeron borcu, kesilecek fatura ve risk hesabına girmez), banka hesapları 0 bakiyeyle açılır.
  * Öz araç yalnız eski panelin "Araçlar" listesindeki plakalardır. Diğer bütün plakalar kiralık (taşeron) araçtır;
    taşeronu eski panelde yazılmamışsa "Taşeronu Belli Olmayan Araçlar" tedarikçisine bağlanır.
  * İl, ilçe ve yükleme/teslim yerlerindeki yer adı yazım hataları düzeltilir (ANTALAYA → ANTALYA); her düzeltme rapora yazılır.
"""
import collections, datetime as dt, html.parser, re, sys
from decimal import Decimal as D
from pathlib import Path
import openpyxl

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else sys.exit(__doc__)
EXP, OUT = SRC / 'exports', SRC / 'aktar'
OUT.mkdir(exist_ok=True)
report = []
note = report.append

# Panelin aktarım şablonlarındaki başlıklar (server/…/ImportService.cs Columns ile aynı).
COLS = {
    'suppliers': ['Ünvan', 'Tür', 'VKN/TCKN', 'Vergi Dairesi', 'Telefon', 'E-posta', 'IBAN', 'Adres', 'İl', 'İlçe', 'Yetkili', 'Vade Gün', 'Devir Borcu', 'Devir Tarihi'],
    'customers': ['Ünvan', 'VKN/TCKN', 'Vergi Dairesi', 'Telefon', 'E-posta', 'Adres', 'Not', 'Devir Bakiyesi', 'Devir Tarihi',
                  'İl', 'İlçe', 'Yetkili', 'e-Fatura Mükellefi', 'PK Etiketi', 'Vade Gün'],
    'drivers': ['Ad Soyad', 'Telefon', 'TC Kimlik No', 'Ehliyet Sınıfı', 'Ehliyet Bitiş', 'SRC Bitiş', 'Psikoteknik Bitiş', 'Tedarikçi'],
    'vehicles': ['Plaka', 'Araç Tipi', 'Marka', 'Model', 'Model Yılı', 'Km', 'Son Bakım', 'Sonraki Bakım', 'Muayene Bitiş', 'Sigorta Bitiş',
                 'Sahiplik', 'Araç Sahibi', 'Dorse Plakası'],
    'trips': ['Yükleme Tarihi', 'Müşteri', 'Plaka', 'Şoför', 'Yükleme İli', 'Yükleme Adresi', 'Teslim İli', 'Teslim Adresi', 'Teslim Tarihi',
              'Yük Cinsi', 'Müşteri Ref No', 'Araç Maliyeti', 'Satış Fiyatı', 'Durum', 'Açıklama', 'Eski Kayıt'],
    'payments': ['Tarih', 'Müşteri', 'Tutar', 'Yöntem', 'Fatura No', 'Çek/Senet No', 'Banka', 'Çek Vadesi', 'Açıklama'],
    'supplier-payments': ['Tarih', 'Tedarikçi', 'Tutar', 'Yöntem', 'Açıklama'],
    'expenses': ['Tarih', 'Kategori', 'Tutar', 'Plaka', 'Şoför', 'Tedarikçi', 'Vadeli', 'Litre', 'Km', 'Açıklama'],
    'cash-accounts': ['Hesap Adı', 'Tür', 'IBAN', 'Devir Bakiyesi', 'Devir Tarihi'],
    'staff': ['Ad Soyad', 'TC Kimlik No', 'Telefon', 'İşe Başlangıç', 'Maaş', 'Not'],
}


class Table(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(); self.rows, self.row, self.cell = [], None, None
    def handle_starttag(self, t, a):
        if t == 'tr': self.row = []
        if t in ('td', 'th') and self.row is not None: self.cell = ''
    def handle_endtag(self, t):
        if t in ('td', 'th') and self.cell is not None: self.row.append(' '.join(self.cell.split())); self.cell = None
        if t == 'tr' and self.row is not None: self.rows.append(self.row); self.row = None
    def handle_data(self, d):
        if self.cell is not None: self.cell += d


def html_table(path, min_cols=3, numbered=False):
    """HTML tabloyu sözlük listesine çevirir. numbered: ilk sütunu sıra no olmayan satırları ("Toplam" satırı) atar."""
    p = Table(); p.feed(Path(path).read_text(encoding='utf-8', errors='replace'))
    rows = [r for r in p.rows if len(r) >= min_cols]
    head = [h.rstrip('↕▼▲ ').strip() for h in rows[0]]
    body = [r for r in rows[1:] if not numbered or r[0].strip().isdigit()]
    return [dict(zip(head, r)) for r in body]


_TR = str.maketrans('iı', 'İI')
def key(s): return ' '.join(str(s or '').translate(_TR).upper().split())
_PLATE = re.compile(r'(\d{2})\s*([A-ZÇĞİÖŞÜ]{1,3})\s*(\d{2,5})')
_PAIR = re.compile(r'^(\d{2}[A-ZÇĞİÖŞÜ]{1,3}\d{2,4})(\d{2}[A-ZÇĞİÖŞÜ]{1,3}\d{2,4})$')
def split_plate(raw):
    """Eski panel çekici ve dorseyi tek alanda tutar ("34ABC123-34DEF456"). (araç, dorse) olarak ayırır, "34 ABC 123" biçimine getirir."""
    s = key(raw)
    if not s: return None, None
    compact = re.sub(r'[^0-9A-ZÇĞİÖŞÜ]', '', s)
    pair = _PAIR.match(compact)
    parts = [pair.group(1), pair.group(2)] if pair and not re.search(r'[-/ ]', s.strip()) else re.split(r'\s*[-/]\s*', s)
    fmt = []
    for p in parts:
        m = _PLATE.fullmatch(re.sub(r'\s+', '', p))
        if m: fmt.append(f'{m.group(1)} {m.group(2)} {m.group(3)}')
    return (fmt[0] if fmt else None), (fmt[1] if len(fmt) > 1 else None)


def clean(s):
    s = ' '.join(str(s or '').split())
    return None if s in ('', '-', '0', 'None') else s
def money(v):
    if v is None or v == '': return D(0)
    if isinstance(v, (int, float)): return D(str(v)).quantize(D('0.01'))
    s = str(v).replace('₺', '').replace('TL', '').strip()
    if not s: return D(0)
    if ',' in s: s = s.replace('.', '').replace(',', '.')
    return D(s).quantize(D('0.01'))


# ---------- yazım düzeltme ----------
# Doğru yazılışlar: 81 il (Cities.cs) + sık geçen ilçe ve semtler. Kamuya açık yer adlarıdır, müşteri verisi değildir.
PROVINCES = re.findall(r'"([^"]+)"', (Path(__file__).resolve().parents[2] / 'server/YesLojistik.Core/Domain/Cities.cs')
                       .read_text(encoding='utf-8').split('All =')[1].split('];')[0])
PLACES = PROVINCES + '''Afyon Antep Urfa Maraş İzmit İskenderun Adapazarı
Adalar Arnavutköy Ataşehir Avcılar Bağcılar Bahçelievler Bakırköy Başakşehir Bayrampaşa Beşiktaş Beykoz Beylikdüzü Beyoğlu
Büyükçekmece Çatalca Çekmeköy Esenler Esenyurt Eyüpsultan Fatih Gaziosmanpaşa Güngören Kadıköy Kağıthane Kartal Küçükçekmece
Maltepe Pendik Sancaktepe Sarıyer Silivri Sultanbeyli Sultangazi Şile Şişli Tuzla Ümraniye Üsküdar Zeytinburnu
Ataköy Bostancı Kozyatağı Göztepe Fulya Etiler Levent Maslak Mecidiyeköy Okmeydanı Poyrazköy Ferhatpaşa Paşaköy Kasımpaşa
Unkapanı Florya İkitelli İstinye Göktürk Ömerli Orhanlı Piyalepaşa Ulus Balat Dudullu Hadımköy Yenibosna Kurtköy Gebze Dilovası
Çayırova Darıca Körfez Gölcük Kartepe Başiskele Derince İzmit
Altındağ Çankaya Etimesgut Gölbaşı Keçiören Mamak Pursaklar Sincan Yenimahalle Polatlı Kızılay Beykent Ostim
Balçova Bayraklı Bornova Buca Çiğli Gaziemir Güzelbahçe Karabağlar Karşıyaka Konak Menemen Narlıdere Torbalı Urla Ödemiş
Selçuk Çeşme Menderes Kemalpaşa Aliağa Alsancak Işıkkent
Muratpaşa Konyaaltı Kepez Döşemealtı Aksu Alanya Manavgat Kemer Serik Kaş Kumluca
Osmangazi Nilüfer Yıldırım Mudanya Gemlik İnegöl Gürsu Kestel Balat Görükle
Bodrum Fethiye Marmaris Datça Milas Menteşe Dalaman Köyceğiz Yalıkavak Ortaca
Edremit Bandırma Ayvalık Gönen Burhaniye Altıeylül Karesi
Pamukkale Merkezefendi Seyhan Çukurova Sarıçam Yüreğir Ceyhan Şahinbey Şehitkamil
Melikgazi Kocasinan Talas Mimarsinan Tepebaşı Odunpazarı Serdivan Arifiye Erenler Akyazı
Çorlu Çerkezköy Süleymanpaşa Ergene Lüleburgaz Antakya Arsuz Dörtyol Defne Mezitli Yenişehir Tarsus Toroslar
Efeler Kuşadası Didim Nazilli Selçuklu Karatay Meram Ereğli Ayvacık Biga Gelibolu Turgutlu Akhisar Yunusemre Şehzadeler
Tatvan Güroymak Elazığ Bulancak'''.split()

_FOLD = str.maketrans('İIÖÜŞÇĞÂÎÛ', 'IIOUSCGAIU')
def fold(s): return key(s).translate(_FOLD)
PLACE_BY_FOLD = {fold(p): p for p in PLACES}

def lev(a, b, limit):
    """İki kelime arasındaki harf farkı (ekleme/silme/değiştirme); limit aşılınca limit+1 döner."""
    if abs(len(a) - len(b)) > limit: return limit + 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1): cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        if min(cur) > limit: return limit + 1
        prev = cur
    return prev[-1]

def place(word):
    """Kelime bir yer adının hatalı yazılışıysa doğru yazılışı (yoksa None). Kısa kelimelere dokunulmaz."""
    f = fold(word)
    if len(f) < 4: return None
    if f in PLACE_BY_FOLD: return PLACE_BY_FOLD[f]
    if len(f) < 5: return None
    limit = 1 if len(f) < 8 else 2
    dist = {p: d for k, p in PLACE_BY_FOLD.items() if (d := lev(f, k, limit)) <= limit}
    best = [p for p, d in dist.items() if d == min(dist.values())]
    return best[0] if len(best) == 1 else None  # en yakın tek aday (ANTALAYA → Antalya; Antakya 2 harf uzakta)

fixes = collections.Counter()
_WORD = r'[A-Za-zÇĞİÖŞÜÂÎÛçğıöşüâîû]+'
def fix_places(text):
    """Serbest metindeki yer adlarını düzeltir ("ANTALAYA 2 YER" → "ANTALYA 2 YER", "KONYA ALTI" → "KONYAALTI"). Büyük harf korunur."""
    if not text: return text
    def same_case(src, right): return key(right) if src.isupper() else right
    parts = re.split(f'({_WORD})', text)  # tek sıralar kelime, çift sıralar aradaki işaretler
    for i in range(1, len(parts) - 2, 2):
        if parts[i] and parts[i + 1] == ' ' and (p :=PLACE_BY_FOLD.get(fold(parts[i] + parts[i + 2]))):
            parts[i], parts[i + 1], parts[i + 2] = same_case(parts[i], p), '', ''
    for i in range(1, len(parts), 2):
        if parts[i] and (p := place(parts[i])): parts[i] = same_case(parts[i], p)
    out = ''.join(parts)
    if out != text: fixes[f'{text} → {out}'] += 1
    return out

def fix_city(v):
    """İl alanı: listedeki resmi yazılış ("SAKARAYA" → "Sakarya"); tanınmazsa olduğu gibi kalır."""
    v = clean(v)
    p = place(v) if v else None
    out = p if p in PROVINCES else ({'Afyon': 'Afyonkarahisar', 'Antep': 'Gaziantep', 'Urfa': 'Şanlıurfa', 'Maraş': 'Kahramanmaraş', 'İzmit': 'Kocaeli'}.get(p) or v)
    if v and out != v and key(out) != key(v): fixes[f'{v} → {out}'] += 1
    return out

def fix_district(v):
    v = clean(v)
    p = place(v) if v and ' ' not in v else None
    if p and key(p) != key(v): fixes[f'{v} → {p}'] += 1
    return p or v

def fix_title(v):
    """Ünvan: noktadan sonra boşluk ("TİC.LTD.ŞTİ." → "TİC. LTD. ŞTİ."); "A.Ş." gibi tek harfli kısaltmalara dokunulmaz."""
    v = clean(v)
    return ' '.join(re.sub(rf'\.(?={_WORD[:-1]}{{2,}})', '. ', v).split()) if v else v


def write(name, entity, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Veri'
    cols = COLS[entity]; ws.append(cols)
    for r in rows: ws.append([r.get(c) for c in cols])
    wb.save(OUT / name)
    note(f'{name}: {len(rows)} satır')


# ---------- girdiler ----------
firms = html_table(EXP / 'firmalar.html')
suppliers_src = html_table(EXP / 'tedarikciler.html')
drivers_src = html_table(EXP / 'soforler.html')
own_src = html_table(EXP / 'araclar.html', min_cols=5) if (EXP / 'araclar.html').exists() else []
expenses_src = html_table(EXP / 'giderler.html', min_cols=5) if (EXP / 'giderler.html').exists() else []
fuel_src = []
if (EXP / 'mazotlar.xlsx').exists():
    _f = list(openpyxl.load_workbook(EXP / 'mazotlar.xlsx', read_only=True).worksheets[0].iter_rows(values_only=True))
    fuel_src = [dict(zip(_f[0], r)) for r in _f[1:] if r and r[0]]
banks_src = html_table(EXP / 'bankalar.html', min_cols=5) if (EXP / 'bankalar.html').exists() else []
staff_src = [r for r in html_table(EXP / 'personeller.html', min_cols=8) if str(r.get('Sıra No', '')).strip().isdigit()] \
    if (EXP / 'personeller.html').exists() else []

ws = openpyxl.load_workbook(EXP / 'sevkiyatlar.xlsx', read_only=True).worksheets[0]
raw = list(ws.iter_rows(values_only=True)); IX = {h: i for i, h in enumerate(raw[0])}
trips_src = [dict((h, r[i]) for h, i in IX.items()) for r in raw[1:]
             if r and isinstance(r[0], str) and re.match(r'\d{4}-\d\d-\d\d$', r[0])]
note(f'Kaynak: {len(firms)} firma, {len(suppliers_src)} tedarikçi, {len(drivers_src)} şoför, {len(trips_src)} sefer, {len(own_src)} öz araç')

own_titles = collections.Counter(key(t['Fatura Başlığı']) for t in trips_src if t['Sevkiyat Durum'] == 'Öz Araç')
OWN = own_titles.most_common(1)[0][0] if own_titles else None  # öz araç seferlerindeki fatura başlığı = firmanın kendisi

# ---------- tedarikçiler ----------
suppliers = {}  # key(unvan) -> satır
for s in suppliers_src:
    title = clean(s.get('Ünvan'))
    if not title: continue
    il_ilce = (s.get('İl / İlçe') or '').split('/')
    suppliers[key(title)] = {
        'Ünvan': fix_title(title), 'Tür': 'Taşeron', 'VKN/TCKN': clean(s.get('VKN/TCKN')), 'Vergi Dairesi': clean(s.get('Vergi Dairesi')),
        'Telefon': clean(s.get('Telefon')), 'E-posta': clean(s.get('E-Posta')), 'IBAN': clean(s.get('IBAN')),
        'İl': fix_city(il_ilce[0]) if il_ilce else None, 'İlçe': fix_district(il_ilce[1]) if len(il_ilce) > 1 else None,
        'Yetkili': clean(s.get('Yetkili')), 'Vade Gün': 30, 'Devir Borcu': 0,
    }
added = set()
for title in [t['Fatura Başlığı'] for t in trips_src if t['Sevkiyat Durum'] == 'Piyasa'] + [d.get('Fatura Başlığı') for d in drivers_src]:
    k = key(title)
    if k and k != OWN and k not in suppliers:
        suppliers[k] = {'Ünvan': fix_title(title), 'Tür': 'Taşeron', 'Vade Gün': 30, 'Devir Borcu': 0}; added.add(k)
if added: note(f'Tedarikçi listesinde olmayıp seferde/şoförde geçen {len(added)} taşeron eklendi (iletişim bilgileri boş).')

# ---------- müşteriler ----------
customers = {}
for f in firms:
    title = clean(f.get('Firma'))
    if not title: continue
    row = {'Ünvan': fix_title(title), 'VKN/TCKN': clean(f.get('VKN/TCKN')), 'Vergi Dairesi': clean(f.get('Vergi Dairesi')),
           'Telefon': clean(f.get('Telefon')), 'E-posta': clean(f.get('E-Posta')), 'İl': fix_city(f.get('İl')), 'İlçe': fix_district(f.get('İlçe')),
           'Yetkili': clean(f.get('Yetkili')), 'Devir Bakiyesi': 0, 'e-Fatura Mükellefi': 'Hayır'}
    customers[key(title)] = row
for t in trips_src:
    k = key(t['Firma Ünvanı'])
    if k and k not in customers:
        customers[k] = {'Ünvan': fix_title(t['Firma Ünvanı']), 'Devir Bakiyesi': 0, 'e-Fatura Mükellefi': 'Hayır'}
        note('Firma listesinde olmayan sefer müşterisi eklendi: 1 kayıt')

# ---------- şoförler ----------
drivers = {}
for d in drivers_src:
    name = clean(d.get('Şöför'))
    if not name: continue
    sup = key(d.get('Fatura Başlığı'))
    drivers[key(name)] = {'Ad Soyad': name, 'Telefon': clean(d.get('GSM')) or clean(d.get('Telefon')), 'TC Kimlik No': clean(d.get('TC')),
                          'Ehliyet Sınıfı': clean(d.get('Ehliyet')), 'Tedarikçi': suppliers[sup]['Ünvan'] if sup in suppliers else None,
                          '_plate': clean(d.get('Plaka'))}
for t in trips_src:
    k = key(t['Şöför İsim'])
    if k and k not in drivers:
        sup = key(t['Fatura Başlığı'])
        drivers[k] = {'Ad Soyad': clean(t['Şöför İsim']), 'Tedarikçi': suppliers[sup]['Ünvan'] if sup in suppliers else None}

# ---------- araçlar ----------
def tr_date(v):
    """'14.04.2027' ya da 'HDI | 14.04.2027' → tarih."""
    m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', str(v or ''))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None

# Öz araç = eski panelin "Araçlar" listesi. Seferdeki "Öz Araç" işareti ya da boş taşeron alanı aracı öz araç yapmaz.
if not own_src: sys.exit('DUR: eski panelin "Araçlar" listesi (araclar.html) yok; öz araçlar ayırt edilemez. extract.mjs ile yeniden çekin.')
OWN_PLATES = {p for a in own_src if (p := split_plate(a.get('Plaka'))[0])}
UNKNOWN_OWNER = 'Taşeronu Belli Olmayan Araçlar'

plate_type = collections.defaultdict(collections.Counter)
plate_owner = collections.defaultdict(collections.Counter)
plate_trailer = collections.defaultdict(collections.Counter)
own_flag_mismatch = set()
bad_plates = set()
for t in trips_src:
    p, trailer = split_plate(t['Plaka'])
    if not p:
        if key(t['Plaka']): bad_plates.add(key(t['Plaka']))
        continue
    t['_plate'] = p
    if trailer: plate_trailer[p][trailer] += 1
    if t['Araç Cinsi']: plate_type[p][clean(t['Araç Cinsi'])] += 1
    if t['Sevkiyat Durum'] == 'Öz Araç' and p not in OWN_PLATES: own_flag_mismatch.add(p)
    if key(t['Fatura Başlığı']) in suppliers: plate_owner[p][key(t['Fatura Başlığı'])] += 1
driver_plates = set()
for d in drivers.values():
    p, trailer = split_plate(d.get('_plate'))
    if d.get('_plate') and not p: bad_plates.add(key(d['_plate']))
    if not p: continue
    driver_plates.add(p)
    if trailer: plate_trailer[p][trailer] += 1
    if d.get('Tedarikçi'): plate_owner[p][key(d['Tedarikçi'])] += 0.5
vehicles = {}
for p in sorted(set(plate_type) | set(plate_owner) | driver_plates | OWN_PLATES):
    typ = plate_type[p].most_common(1)[0][0] if plate_type[p] else 'Kamyon'
    row = {'Plaka': p, 'Araç Tipi': typ, 'Dorse Plakası': plate_trailer[p].most_common(1)[0][0] if plate_trailer[p] else None}
    if p in OWN_PLATES: row['Sahiplik'] = 'Özmal'
    elif plate_owner[p]: row.update({'Sahiplik': 'Kiralık', 'Araç Sahibi': suppliers[plate_owner[p].most_common(1)[0][0]]['Ünvan']})
    else:
        suppliers.setdefault(key(UNKNOWN_OWNER), {'Ünvan': UNKNOWN_OWNER, 'Tür': 'Taşeron', 'Vade Gün': 30, 'Devir Borcu': 0})
        row.update({'Sahiplik': 'Kiralık', 'Araç Sahibi': UNKNOWN_OWNER})
    vehicles[p] = row

for a in own_src:  # öz araçların marka, sigorta ve muayene bilgileri
    p, _ = split_plate(a.get('Plaka'))
    if not p: continue
    row = vehicles[p]
    mm = clean(a.get('Marka Model')) or ''
    year = re.search(r'\b(19|20)\d{2}\b', mm)
    row.update({'Marka': clean(re.sub(r'\b(19|20)\d{2}\b', '', mm)), 'Model Yılı': int(year.group(0)) if year else None,
                'Muayene Bitiş': tr_date(a.get('Muay. / Tarih')), 'Sigorta Bitiş': tr_date(a.get('Sigorta / Tarih'))})
    if clean(a.get('Tip')): row['Araç Tipi'] = clean(a.get('Tip')).title()
unknown = sum(v.get('Araç Sahibi') == UNKNOWN_OWNER for v in vehicles.values())
note(f"Araçlar: {len(OWN_PLATES)} öz araç (eski panelin Araçlar listesi"
     + (f'; {len(own_src) - len(OWN_PLATES)} plaka orada iki kez yazılmış' if len(own_src) > len(OWN_PLATES) else '') + f"), {len(vehicles) - len(OWN_PLATES)} taşeron aracı"
     + (f'; {unknown} aracın taşeronu eski panelde yazılmamış, "{UNKNOWN_OWNER}" altına alındı (araç kartından düzeltilebilir)' if unknown else '') + '.')
if own_flag_mismatch:
    note(f'  {len(own_flag_mismatch)} plaka eski panelde "Öz Araç" seferinde geçiyor ama Araçlar listesinde yok; taşeron aracı sayıldı: '
         + ', '.join(sorted(own_flag_mismatch)))
for d in drivers.values(): d.pop('_plate', None)
if bad_plates: note(f'{len(bad_plates)} plaka biçimi tanınamadı; bu plakalı seferler alınmadı (rapora bakın).')

# ---------- seferler ----------
# Hepsi eski kayıt: geçmişte ve raporlarda görünür, borç/alacak, kesilecek fatura ve risk hesabına girmez.
trips, skipped = [], 0
for t in trips_src:
    cust, plate, drv = key(t['Firma Ünvanı']), t.get('_plate'), key(t['Şöför İsim'])
    if not (cust and plate and drv):
        skipped += 1; continue
    invoice = clean(t['Kesilen Fatura'])
    extras = [f"Sevkiyat {t['Sevkiyat No']}" if t.get('Sevkiyat No') else None,
              clean(t.get('Açıklama')),
              f"Komisyon {money(t['Komisyon'])}" if money(t['Komisyon']) else None,
              f"Masraf {money(t['Masraf'])}" if money(t['Masraf']) else None,
              f"Şoför primi {money(t['Şöför Prim'])}" if money(t['Şöför Prim']) else None,
              f"Teslim evrak {clean(t['Teslim Evrak No'])}" if clean(t.get('Teslim Evrak No')) else None,
              f"Fatura {invoice}" if invoice not in (None, 'Komisyon İşi') else None]
    d = dt.date.fromisoformat(t['Tarih'])
    trips.append({'Yükleme Tarihi': d, 'Müşteri': customers[cust]['Ünvan'], 'Plaka': plate, 'Şoför': drivers[drv]['Ad Soyad'],
                  'Yükleme Adresi': fix_places(clean(t['Yükleme Noktası'])) or '-', 'Teslim Adresi': fix_places(clean(t['İndirme Noktası'])) or '-',
                  'Teslim Tarihi': d,
                  'Yük Cinsi': clean(t['Ürün']), 'Araç Maliyeti': float(money(t['Şöför Fiyat'])), 'Satış Fiyatı': float(money(t['Müşteri Fiyat'])),
                  'Durum': 'Teslim Edildi', 'Açıklama': ' · '.join(x for x in extras if x)[:1000], 'Eski Kayıt': 'Evet'})
if skipped: note(f'{skipped} sefer satırı müşteri/plaka/şoför boş olduğu için alınmadı.')
note(f'Seferler: {len(trips)} sefer, hepsi eski kayıt (borç/alacak doğurmaz).')
note('Cari devri yok: müşteri ve tedarikçiler 0 bakiyeyle açılır.')

# ---------- giderler ve mazot ----------
# Eski "Giderler" listesi bir kasa çıkış defteri: taşeron ödemeleri ("Fatura No …", "VKN …"), NAKLİYESPOTARAÇLAR ve
# mahsuplaşmalar gider değil, cari hareketidir; borç/alacak taşınmadığı için alınmaz. Gerçek giderler alınır.
PAYMENT_CATS = {'NAKLİYESPOTARAÇLAR', 'MAHSUPLAŞMA'}
def expense_category(cat, text):
    t = key(f'{cat} {text}')
    if 'HGS' in t or 'GEMİ' in t or 'FERİBOT' in t or 'OTOYOL' in t or 'KÖPRÜ' in t: return 'Otoyol'
    if 'YAKIT' in t or 'MAZOT' in t: return 'Yakıt'
    if 'HARÇLIK' in t: return 'Harcırah'
    if 'TAMİR' in t or 'BAKIM' in t or 'SERVİS' in t: return 'Bakım'
    if 'LASTİK' in t: return 'Lastik'
    if 'SİGORTA' in t or 'KASKO' in t: return 'Sigorta'
    if 'VERGİ' in t or 'SSK' in t or 'SGK' in t or 'MTV' in t: return 'Vergi'
    return 'Diğer'
expenses, skipped_exp = [], collections.Counter()
for e in expenses_src:
    cat, what, note_ = key(e.get('Kategori')), clean(e.get('Gider')), clean(e.get('Not')) or ''
    amount = money(e.get('Tutar'))
    if (cat in PAYMENT_CATS or note_.startswith(('Fatura No', 'VKN')) or (not cat and (not what or key(what) == 'KISMI ÖDEME'))) or amount <= 0:
        skipped_exp[cat or 'kategorisiz ödeme'] += amount; continue
    plate, _ = split_plate(cat)
    known = plate in vehicles  # bilinmeyen plaka için araç açılmaz (öz araç sanılmasın); plaka açıklamada kalır
    expenses.append({'Tarih': tr_date(e.get('Tarih')), 'Kategori': expense_category('' if plate else cat, what or ''),
                     'Tutar': float(amount), 'Plaka': plate if known else None,
                     'Açıklama': ' · '.join(x for x in [None if known else clean(e.get('Kategori')), what, note_ or None] if x)[:500] or None})
for f in fuel_src:
    plate, _ = split_plate(f.get('Plaka'))
    if not plate: continue
    known = plate in vehicles
    expenses.append({'Tarih': tr_date(f.get('Tarih')), 'Kategori': 'Yakıt', 'Tutar': float(money(f.get('Tutar'))), 'Plaka': plate if known else None,
                     'Litre': float(money(f.get('Litre'))) or None, 'Km': int(f['Yeni KM']) if known and f.get('Yeni KM') else None,
                     'Açıklama': ' · '.join(x for x in ['Mazot', None if known else plate, clean(f.get('Petrol'))] if x)})
# Eski panelde aynı gün aynı tutarda iki ayrı gider olabilir; aktarım bunları tekrar sayıp atlamasın diye sıra eklenir.
seen = collections.Counter()
for e in expenses:
    k = (e['Tarih'], e['Kategori'], e['Tutar'], e.get('Plaka'), e.get('Açıklama'))
    seen[k] += 1
    if seen[k] > 1: e['Açıklama'] = f"{e.get('Açıklama') or ''} ({seen[k]}. kayıt)".strip()
if expenses_src or fuel_src:
    note(f'Giderler: {len(expenses)} kayıt (mazot {len(fuel_src)}). Cari hareketi olan ödemeler gider sayılmadı: '
         + ', '.join(f'{k} {v:,.2f} TL' for k, v in skipped_exp.most_common()))

# ---------- banka hesapları ----------
cash_accounts = [{'Hesap Adı': clean(b.get('Hesap İsmi')), 'Tür': 'Banka', 'Devir Bakiyesi': 0}
                 for b in banks_src if clean(b.get('Hesap İsmi'))]
if cash_accounts: note(f'Banka hesapları: {len(cash_accounts)} hesap, 0 bakiyeyle açılır.')

# ---------- personel ----------
staff = [{'Ad Soyad': clean(p.get('İsim / Soyisim')), 'TC Kimlik No': clean(p.get('Tc')), 'Telefon': clean(p.get('Telefon')),
          'İşe Başlangıç': tr_date(p.get('Başlangıç Tarih')), 'Maaş': float(money(p.get('Maaş'))), 'Not': clean(p.get('Not'))}
         for p in staff_src if clean(p.get('İsim / Soyisim'))]
if staff: note(f'Personel: {len(staff)} kişi aylık maaşıyla.')

write('1-tedarikciler.xlsx', 'suppliers', list(suppliers.values()))
write('2-musteriler.xlsx', 'customers', list(customers.values()))
write('3-soforler.xlsx', 'drivers', list(drivers.values()))
write('4-araclar.xlsx', 'vehicles', list(vehicles.values()))
write('5-seferler.xlsx', 'trips', trips)
if expenses: write('8-giderler.xlsx', 'expenses', expenses)
if cash_accounts: write('9-banka-hesaplari.xlsx', 'cash-accounts', cash_accounts)
if staff: write('10-personeller.xlsx', 'staff', staff)
if fixes:
    note(''); note(f'YAZIM DÜZELTMELERİ ({sum(fixes.values())} alan):')
    for k, n in sorted(fixes.items()): note(f'  {k}' + (f'  ({n} kez)' if n > 1 else ''))
(OUT / 'rapor.txt').write_text('\n'.join(report) + '\n', encoding='utf-8')
print('\n'.join(report))
